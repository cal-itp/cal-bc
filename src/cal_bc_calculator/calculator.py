import pathlib
from functools import cached_property
from io import BytesIO
from itertools import groupby

import formualizer
import openpyxl


class CellResolver:
    def __init__(self, workbook: formualizer.Workbook) -> None:
        self.workbook = workbook

    @property
    def defined_names(self) -> dict[str, tuple[str, int, int]]:
        return {
            nr['name']: (nr["sheet"], nr["start_row"], nr["start_col"])
            for nr in self.workbook.get_named_ranges()
        }

    def excel_column_index(self, column):
        n = 0
        for char in column:
            n = n * 26 + 1 + ord(char) - ord('A')
        return n

    def parse(self, address):
        sheet, cell = address.split("!")
        column, row = ["".join(g) for _, g in groupby(cell, str.isalpha)]
        return (sheet, int(row), self.excel_column_index(column))

    def resolve(self, address: str) -> tuple[str, int, int]:
        if address in self.defined_names:
            return self.defined_names[address]
        else:
            return self.parse(address)


class BytesCalculator:
    def __init__(self, input: bytes) -> None:
        self.input = input
        self._values = {}

    @cached_property
    def workbook(self) -> formualizer.Workbook:
        return formualizer.load_workbook_bytes(self.input, backend="umya")

    @property
    def resolver(self) -> CellResolver:
        return CellResolver(workbook=self.workbook)

    def save(self, output_file: str) -> None:
        pathlib.Path(output_file).write_bytes(self.to_bytes())

    def to_bytes(self) -> bytes:
        wb = openpyxl.load_workbook(filename=BytesIO(self.input), keep_vba=True)
        for address, value in self._values.items():
            if "!" not in address:
                sheet, cell = next(iter(wb.defined_names[address].destinations))
            else:
                sheet, cell = address.split("!")
            wb[sheet][cell] = value
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.read()

    def write(self, cell_values: dict[str, any]) -> None:
        for key, value in cell_values.items():
            self._values[key] = value

    def evaluate(self, cells: list[str]) -> list[any]:
        workbook = formualizer.load_workbook_bytes(self.to_bytes(), backend="umya")
        resolver = CellResolver(workbook=workbook)
        return workbook.evaluate_cells([resolver.resolve(cell) for cell in cells])


class Calculator:
    def __init__(self, input: str) -> None:
        self.input = input

    @cached_property
    def bytes_calculator(self) -> BytesCalculator:
        workbook_bytes = pathlib.Path(self.input).read_bytes()
        return BytesCalculator(workbook_bytes)

    def save(self, output_file: str) -> None:
        self.bytes_calculator.save(output_file=output_file)

    def to_bytes(self) -> bytes:
        return self.bytes_calculator.to_bytes()

    def write(self, cell_values: dict[str, any]) -> None:
        self.bytes_calculator.write(cell_values=cell_values)

    def evaluate(self, cells: list[str]) -> list[any]:
        return self.bytes_calculator.evaluate(cells=cells)
